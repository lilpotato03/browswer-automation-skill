"""Write leads out as Excel (default), CSV or JSON."""
from __future__ import annotations

import csv
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict, List

META_COLUMNS = ["site", "keywords", "first_seen_at", "last_seen_at", "times_seen"]
META_LABELS = {
    "site": "Site", "keywords": "Keyword(s)", "first_seen_at": "First seen",
    "last_seen_at": "Last seen", "times_seen": "Times seen",
}


def _columns(leads: List[dict], sites: Dict[str, object], wanted: List[str]) -> List[str]:
    cols: List[str] = list(wanted)
    for mod in sites.values():
        for c in getattr(mod, "COLUMNS", []):
            if c not in cols:
                cols.append(c)
    for lead in leads:
        for c in lead["data"]:
            if c not in cols:
                cols.append(c)
    return cols


def _flatten(lead: dict, cols: List[str]) -> List[str]:
    return [lead["data"].get(c, "") for c in cols] + [lead[m] for m in META_COLUMNS]


def _header(cols: List[str]) -> List[str]:
    return [c.replace("_", " ").title() for c in cols] + [META_LABELS[m] for m in META_COLUMNS]


def _target(folder: Path, stem: str, ext: str) -> Path:
    """Never overwrite an earlier run: a second run the same day gets a time suffix."""
    path = folder / f"{stem}.{ext}"
    return path if not path.exists() else folder / f"{stem}_{datetime.now():%H%M%S}.{ext}"


def _publish(paths: List[Path], folder: Path, out: dict) -> None:
    """latest.<ext> in the project's output folder + a copy into every copy_to folder."""
    for p in paths:
        tag = p.stem.rsplit("_", 1)[-1] if p.suffix == ".csv" else ""
        if out.get("write_latest", True):
            latest = folder / (f"latest_{tag}{p.suffix}" if tag in ("new", "all") else f"latest{p.suffix}")
            try:
                shutil.copy2(p, latest)
            except PermissionError:  # latest.xlsx open in Excel (Windows): skip, the dated file exists
                pass
        for dest in out.get("copy_to", []):
            d = Path(dest).expanduser()
            d.mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, d / p.name)


def export(db, run_id: int, settings: dict, sites: Dict[str, object], root: Path) -> List[Path]:
    out = settings.get("output", {})
    fmt = out.get("format", "xlsx").lower()
    folder = root / out.get("folder", "output")
    folder.mkdir(parents=True, exist_ok=True)
    stem = out.get("file_name", "{project}_{date}").format(
        date=datetime.now().strftime("%Y-%m-%d"),
        project=settings.get("project_slug") or "leads")

    new = db.leads("first_run_id = ?", (run_id,))
    everything = db.leads() if out.get("include_all_leads_sheet", True) else []
    cols = _columns(new + everything, sites, out.get("columns", []))
    if out.get("only_listed_columns") and out.get("columns"):
        cols = list(out["columns"])

    if fmt == "csv":
        paths = [_target(folder, stem + "_new", "csv")]
        _write_csv(paths[0], cols, new)
        if everything:
            paths.append(_target(folder, stem + "_all", "csv"))
            _write_csv(paths[1], cols, everything)
    elif fmt == "json":
        paths = [_target(folder, stem, "json")]
        paths[0].write_text(json.dumps({"new": [l["data"] for l in new],
                                        "all": [l["data"] for l in everything]},
                                       indent=2, ensure_ascii=False), encoding="utf-8")
    else:
        paths = [_target(folder, stem, "xlsx")]
        _write_xlsx(paths[0], cols, new, everything, db.run_history(), db.errors_for_run(run_id))
    _publish(paths, folder, out)
    return paths


def _write_csv(path: Path, cols: List[str], leads: List[dict]) -> None:
    # utf-8-sig so Excel opens accented characters correctly
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(_header(cols))
        for lead in leads:
            w.writerow(_flatten(lead, cols))


def _write_xlsx(path: Path, cols, new, everything, runs, errors) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")

    def sheet(ws, header, rows):
        ws.append(header)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, start=1):
            longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith(("http://", "https://")):
                    cell.hyperlink = v
                    cell.font = Font(color="0563C1", underline="single")

    ws = wb.active
    ws.title = "New this run"
    sheet(ws, _header(cols), [_flatten(l, cols) for l in new])
    if everything:
        sheet(wb.create_sheet("All leads"), _header(cols), [_flatten(l, cols) for l in everything])
    run_cols = ["id", "started_at", "finished_at", "status", "sites", "keywords",
                "new_leads", "seen_leads", "errors"]
    sheet(wb.create_sheet("Run history"), [c.replace("_", " ").title() for c in run_cols],
          [[r[c] for c in run_cols] for r in runs])
    if errors:
        sheet(wb.create_sheet("Problems"), ["Site", "Keyword", "What went wrong", "Screenshot"],
              [[e["site"], e["keyword"], e["message"][:300], e["screenshot"]] for e in errors])
    wb.save(path)
